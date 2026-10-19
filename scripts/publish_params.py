#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把「来源表格 + 稿件 + 网页清单」合成一份作业清单，供发布与回写使用。

每篇一条记录，字段就是后面两步真正要用的东西：

    number / domain   ← 来源表格（按稿件文件名开头的编号匹配「编号」列）
    title / body      ← 稿件 .txt（与 draft_check / render_pages 同一解析口径）
    site_link         ← 网页清单 _pages.json 里的短 ID + --site-base
    platform_ok       ← 标题长度是否满足各平台上限（头条 30 / 搜狐 72 / CSDN 100）

body 用 \\n 连接各段——正好是 media_publish 需要的形状，也是稿件原文，
**不要手抄正文**，手抄必然与网页、与稿件三处不一致。

用法：

    python publish_params.py --dir "<项目文件夹>" --out jobs.json
    python publish_params.py --dir "<项目文件夹>" --xlsx "<来源表格.xlsx>" \\
        --manifest "<项目文件夹>/article-detail/_pages.json" \\
        --site-base "http://xinwen.wangbo.wang/article-detail/" --out jobs.json

退出码：0 = 全部就绪，1 = 有稿件解析失败或缺料，2 = 用法错误。
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from render_pages import DraftError, discover, parse_draft  # noqa: E402

# 各平台的标题上限（实测自 recipe 的 params.max_len）
PLATFORM_TITLE_LIMIT = {"sohu": 72, "toutiao": 30, "csdn": 100}
NUMBER_RE = re.compile(r"^\D{0,3}(\d{1,4})\D")
DOMAIN_RE = re.compile(r"[\u4e00-\u9fffA-Za-z0-9]{1,30}\.网址")


def norm_no(value) -> str:
    """编号归一：去掉空白，纯数字时去掉前导零（表格里是 1，文件名里常写 01）。"""
    s = "" if value is None else str(value).strip()
    return str(int(s)) if s.isdigit() else s


def norm_text(s: str) -> str:
    """比对用归一：全角括号→半角、去掉所有空白。"""
    return (s.replace("（", "(").replace("）", ")")
             .replace(" ", "").replace("\u3000", "").replace("\n", ""))


def table_index(xlsx: Path) -> dict:
    """从来源表格取 编号 → {编号, 域名, 企业名称, 简称, 行号}。只读。"""
    try:
        import openpyxl
    except Exception:
        print("（没装 openpyxl，跳过表格匹配：number/domain 留空）", file=sys.stderr)
        return {}
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    ws = wb[wb.sheetnames[0]]
    header = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        header[c] = "" if v is None else str(v).strip()

    def col(*keys):
        for c, h in header.items():
            if any(k in h for k in keys):
                return c
        return None

    c_no, c_dom = col("编号"), col("域名")
    c_name, c_short = col("企业名称", "公司名称"), col("简称")
    out = {}
    if c_no is None:
        print("（来源表格里找不到「编号」列，跳过表格匹配）", file=sys.stderr)
        return {}
    for r in range(2, ws.max_row + 1):
        no = ws.cell(r, c_no).value
        if no is None or str(no).strip() == "":
            continue
        key = norm_no(no)
        out[key] = {
            "number": key,
            "row": r,
            "domain": str(ws.cell(r, c_dom).value or "").strip() if c_dom else "",
            "company": str(ws.cell(r, c_name).value or "").strip() if c_name else "",
            "short": str(ws.cell(r, c_short).value or "").strip() if c_short else "",
        }
    wb.close()
    return out


def manifest_map(path: Path | None) -> dict:
    """稿件文件名（茎）→ {id, file}，来自 render_pages.py --short-id 的清单。"""
    if not path or not path.exists():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    out = {}
    for e in data.get("pages", []):
        src = Path(e.get("source", ""))
        if src.name:
            out[src.stem] = {"id": e.get("id") or Path(e["file"]).stem, "file": e["file"]}
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description="合成发布/回写作业清单（表格 + 稿件 + 网页清单）")
    ap.add_argument("--dir", type=Path, required=True, help="稿件目录")
    ap.add_argument("--xlsx", type=Path, help="来源表格（取 编号 / 域名 / 企业名称 / 简称）")
    ap.add_argument("--manifest", type=Path, help="网页清单 _pages.json（取短 ID → 网页链接）")
    ap.add_argument("--kepu-dir", type=Path,
                    help="科普稿目录（默认与 --dir 相同）；按 `NN-科普-*.txt` 与案例稿配对")
    ap.add_argument("--kepu-manifest", type=Path,
                    help="科普稿网页清单 _pages-科普.json（取短 ID → 科普网页链接）")
    ap.add_argument("--allow-missing-kepu", action="store_true",
                    help="允许案例稿没有配对科普稿（默认必须成对：每篇案例稿都要配一篇科普稿）")
    ap.add_argument("--site-base", default="",
                    help="网页所在站点的目录前缀，如 http://xinwen.wangbo.wang/article-detail/")
    ap.add_argument("--out", type=Path, help="输出作业清单 JSON；不给则打印到屏幕")
    a = ap.parse_args()

    drafts = discover(a.dir)
    if not drafts:
        print("没有找到任何稿件（文件名需含「稿件」）", file=sys.stderr)
        return 2

    tindex = table_index(a.xlsx) if a.xlsx else {}
    mindex = manifest_map(a.manifest)
    mindex.update(manifest_map(a.kepu_manifest))          # 科普稿清单也并进来
    kepu_dir = a.kepu_dir or a.dir

    jobs, problems, notes = [], [], []
    base = a.site_base.rstrip("/") + "/" if a.site_base else ""
    for path in drafts:
        try:
            title, paras, truncated = parse_draft(path)
        except DraftError as exc:
            problems.append(f"{path.name}：{exc}")
            continue

        m = NUMBER_RE.match(path.name)
        number = m.group(1) if m else None
        row = tindex.get(norm_no(number) if number else "", {})
        text = title + "\n" + "\n".join(paras)
        text_n = norm_text(text)
        if a.xlsx and not row:
            problems.append(f"{path.name}：文件名开头的编号 {number!r} 在表格里找不到对应行")
        elif row:
            # 交叉核验：只按编号匹配不安全——表传错了会把链接回写到别的企业那一行。
            # 稿件按 H1 必须含企业全称，所以「全称或简称出现在稿件里」才算身份对上；
            # 域名只作辅助信号（表格的「注册域名」可能写中文、正文用的是拼音域名）。
            hits = {k: (row[k] and norm_text(row[k]) in text_n)
                    for k in ("company", "short", "domain")}
            if not any(hits.values()):
                problems.append(
                    f"{path.name}：编号 {number} 匹配到表格第 {row['row']} 行"
                    f"（域名={row['domain'] or '-'}／企业={row['company'] or '-'}），"
                    f"但稿件里既没有这家企业名称也没有这个域名 —— "
                    f"表格与稿件疑似不是同一批，停下人工确认")
            elif not hits["domain"] and row["domain"]:
                found = sorted(set(DOMAIN_RE.findall(text))) or ["（稿件里没写域名）"]
                notes.append(
                    f"{path.name}：表格域名={row['domain']}，稿件里出现的域名={found}，"
                    f"两者不一致 —— 回写按表格行定位（不影响正确性），"
                    f"但交付说明里要写明这个差异，并请用户确认以哪个为准")

        page = mindex.get(path.stem, {})
        site_link = f"{base}{page['id']}.html" if (base and page.get("id")) else ""

        # ---- 配对科普稿：NN-科普-方向K-*.txt --------------------------------
        kepu = None
        if number:
            khits = sorted(kepu_dir.glob(f"{int(norm_no(number)):02d}-科普-*.txt"))
            if len(khits) > 1:
                problems.append(f"{path.name}：配对到 {len(khits)} 篇科普稿 {[h.name for h in khits]}，"
                                f"无法确定是哪一篇")
            elif not khits:
                if not a.allow_missing_kepu:
                    problems.append(
                        f"{path.name}：找不到配对的科普稿（应为 "
                        f"{int(norm_no(number)):02d}-科普-方向K-*.txt）—— "
                        f"每篇案例稿都要配一篇科普稿，缺了不发")
            else:
                kp = khits[0]
                try:
                    ktitle, kparas, _ = parse_draft(kp)
                except DraftError as exc:
                    problems.append(f"{kp.name}：{exc}")
                else:
                    kpage = mindex.get(kp.stem, {})
                    d = re.search(r"方向\s*([1-6])", kp.name)
                    kepu = {
                        "file": kp.name,
                        "source": str(kp),
                        "direction": int(d.group(1)) if d else None,
                        "title": ktitle,
                        "title_len": len(ktitle),
                        "body": "\n".join(kparas),
                        "body_paras": len(kparas),
                        "page_id": kpage.get("id", ""),
                        "site_link": (f"{base}{kpage['id']}.html"
                                      if (base and kpage.get("id")) else ""),
                        "platform_ok": {k: len(ktitle) <= v
                                        for k, v in PLATFORM_TITLE_LIMIT.items()},
                    }

        jobs.append({
            "source": str(path),
            "file": path.name,
            "number": number,
            "domain": row.get("domain", ""),
            "company": row.get("company", ""),
            "title": title,
            "title_len": len(title),
            "body": "\n".join(paras),
            "body_paras": len(paras),
            "page_id": page.get("id", ""),
            "site_link": site_link,
            "platform_ok": {k: len(title) <= v for k, v in PLATFORM_TITLE_LIMIT.items()},
            "kepu": kepu,
            "truncated": truncated,
        })

    print(f"{'案例稿':<26}{'编号':<5}{'标题字数':<9}{'网页ID':<9}"
          f"{'科普稿':<24}{'方向':<5}{'科普网页ID':<11}平台标题适配")
    print("-" * 122)
    for j in jobs:
        ok = " ".join(f"{k}{'✓' if v else '✗'}" for k, v in j["platform_ok"].items())
        k = j.get("kepu")
        print(f"{j['file'][:24]:<26}{str(j['number'] or '-'):<5}{j['title_len']:<9}"
              f"{(j['page_id'] or '-'):<9}"
              f"{(k['file'][:22] if k else '(缺科普稿)'):<24}"
              f"{(('方向' + str(k['direction'])) if (k and k['direction']) else '-'):<5}"
              f"{(k['page_id'] if (k and k['page_id']) else '-'):<11}{ok}")
        if not all(j["platform_ok"].values()):
            print(f"    ↑「{j['title']}」超出某些平台的标题上限；"
                  f"**不要为平台单独改标题**（会与稿件、网页三处不一致），"
                  f"该平台改用另一篇稿件或跳过该平台")
        if k and not all(k["platform_ok"].values()):
            print(f"    ↑ 科普稿「{k['title']}」标题 {k['title_len']} 字超限："
                  f"科普稿要与案例稿发同一平台，必须压到 30 字内（改标题，不要改方向）")
    for n in notes:
        print(f"  注意  {n}")
    for p in problems:
        print("  FAIL  " + p, file=sys.stderr)

    if a.out:
        a.out.write_text(json.dumps(jobs, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"\n作业清单 → {a.out}（{len(jobs)} 条）")

    print(f"\nRESULT: {'PASS' if not problems else 'PARTIAL'} — {len(jobs)} 条就绪"
          f"{'，' + str(len(problems)) + ' 条有问题' if problems else ''}")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
