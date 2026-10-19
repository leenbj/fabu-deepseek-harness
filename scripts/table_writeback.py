# -*- coding: utf-8 -*-
"""来源表格回写工具（授权后使用）——状态机与白名单见 references/table-writeback.md。

设计原则
  1. 先备份，再写；--dry-run 只打印不落盘。
  2. 按行精确定位（默认用「域名」列），逐单元格赋值，绝不整列/整表写入。
  3. 白名单：编号列 / 状态列 / 媒体列（表头=账号代号）/ 末尾两个链接列：
       「发布链接」= 自媒体发布链接（--link）
       「网页链接」= 上传到 FTP 后的网页链接（--site-link）
     两个链接列不存在时各自追加一列，**不覆盖已有的其他列**。
  4. 幂等：已「已发布」的行不被「已写作」降级；重复执行同一状态报告 unchanged；
     同一链接重复写入报告 unchanged，链接已有且不同则拒绝覆盖。
  5. 不确定就不写：定位不唯一、状态列口径冲突、链接冲突 → 停下并报错，不猜不兜底。

用法
  python table_writeback.py --xlsx <表> --domain 某某.网址 --number 76 --status 已写作
  python table_writeback.py --xlsx <表> --domain 某某.网址 --status 已发布 --account a-1 \
      --link <自媒体链接> --site-link <网页链接>
  python table_writeback.py --xlsx <表> --domain 某某.网址 --site-link <网页链接>   # 只补网页链接
  python table_writeback.py --xlsx <表> --batch jobs.json [--dry-run]
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

# 自媒体发布链接列（用户表里通常叫「发布链接」）
LINK_HEADER = "发布链接"
LINK_ALIASES = ("发布链接", "媒体链接", "自媒体链接", "发布地址", "链接")
# 网页链接列（上传到 FTP 之后的页面地址）
SITE_LINK_HEADER = "网页链接"
SITE_LINK_ALIASES = ("网页链接", "网站链接", "网页地址", "站点链接", "FTP链接", "ftp链接")
VALID_STATUS = ("已写作", "已发布")
STATUS_RANK = {"已写作": 1, "已发布": 2}


def out(msg: str = "") -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    print(msg)


class TableError(Exception):
    """表格结构或数据冲突：按契约必须停下报告，不得猜。"""


def _is_site_link(header: str) -> bool:
    return any(k in header for k in ("网页", "网站", "站点", "FTP", "ftp"))


def find_columns(ws) -> dict:
    """按表头关键词定位列。注意：本类表格表头可能写错或残留数据值，
    因此这里只做「定位」，列义核对由调用方打印样例完成。"""
    headers = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        headers[c] = "" if v is None else str(v).strip()

    def find(pred):
        for c in sorted(headers):
            if pred(headers[c]):
                return c
        return None

    cols = {
        "number": find(lambda h: "编号" in h),
        "status": find(lambda h: "状态" in h),
        "domain": find(lambda h: "域名" in h),
        "link": find(lambda h: h in LINK_ALIASES
                     or ("链接" in h and not _is_site_link(h))),
        "site_link": find(lambda h: h in SITE_LINK_ALIASES
                          or ("链接" in h and _is_site_link(h))),
        "headers": headers,
    }
    if cols["domain"] is None:
        raise TableError("找不到「域名」列：无法按域名精确定位行，请人工确认表头")
    return cols


def locate_row(ws, cols: dict, *, domain=None, number=None, row=None) -> int:
    """定位目标行；不唯一、找不到、或多个条件互相矛盾即报错。

    定位优先级与契约一致：**域名优先**；同时给出域名与编号时两者必须指向同一行。
    """
    dom_col = cols["domain"]
    if row is not None:
        if row < 2 or row > ws.max_row:
            raise TableError(f"--row {row} 超出数据区 2..{ws.max_row}")
        got = ws.cell(row, dom_col).value
        if not got:
            raise TableError(f"--row {row} 的域名列为空，拒绝盲写")
        return row

    by_domain = []
    if domain:
        by_domain = [r for r in range(2, ws.max_row + 1)
                     if str(ws.cell(r, dom_col).value or "").strip() == str(domain).strip()]
    by_number = []
    if number is not None:
        if cols["number"] is None:
            raise TableError("找不到「编号」列，无法按编号定位")
        by_number = [r for r in range(2, ws.max_row + 1)
                     if str(ws.cell(r, cols["number"]).value) == str(number)]

    if by_domain and by_number:
        if set(by_domain) != set(by_number):
            raise TableError(f"域名指向 {by_domain}，编号指向 {by_number}，两者矛盾，拒绝写入")
    hits = by_domain or by_number
    if not hits:
        raise TableError(f"定位不到行（domain={domain!r} number={number!r}）——不猜、不模糊匹配")
    if len(hits) > 1:
        raise TableError(f"定位到 {len(hits)} 行 {hits}，不唯一，拒绝写入")
    return hits[0]


def media_col(cols: dict, account: str, explicit: str | None) -> int:
    want = explicit or account
    for c, h in cols["headers"].items():
        if h == want:
            return c
    raise TableError(f"找不到表头为 {want!r} 的媒体列（现有表头：{list(cols['headers'].values())}）")


def ensure_col(ws, cols: dict, key: str, header: str, plan: list, cursor: list) -> int:
    """返回该白名单链接列的列号；不存在则规划在最后一列之后新增一列。

    `cursor` 是只含一个整数的列表，记录下一个可用的新列号——因为
    `ws.max_column` 在保存前不会变，连加两列必须自己往前推。
    """
    if cols.get(key):
        return cols[key]
    new_col = cursor[0]
    cursor[0] += 1
    plan.append((1, new_col, None, header, f"新增「{header}」列表头"))
    cols[key] = new_col
    return new_col


def write_link(ws, row: int, col: int, value: str, what: str, changes: list) -> None:
    old = ws.cell(row, col).value
    old_s = "" if old is None else str(old).strip()
    if old_s == value.strip():
        changes.append((row, col, old, old_s, f"{what}已一致，unchanged"))
    elif old_s:
        raise TableError(f"第 {row} 行{what}列已有值且不同：{old_s!r} ≠ 新值，"
                         f"拒绝覆盖（请人工确认）")
    else:
        changes.append((row, col, old, value.strip(), f"写{what}"))


def plan_one(ws, cols: dict, row: int, job: dict, cursor: list) -> list:
    """算出这一行要改哪些格子，返回 [(row, col, old, new, why)]。"""
    changes = []

    # 1) 状态 + 编号
    status = job.get("status")
    if status:
        if status not in VALID_STATUS:
            raise TableError(f"状态必须是 {VALID_STATUS} 之一，收到 {status!r}")
        if cols["status"] is None:
            raise TableError("找不到「状态」列，拒绝写入（请人工确认表头）")
        sc = cols["status"]
        old = ws.cell(row, sc).value
        old_s = "" if old is None else str(old).strip()
        if old_s in STATUS_RANK:
            if STATUS_RANK[status] < STATUS_RANK[old_s]:
                raise TableError(f"第 {row} 行已是「{old_s}」，拒绝降级为「{status}」（状态只前进不回退）")
            if old_s == status:
                changes.append((row, sc, old, old_s, "状态已是目标值，unchanged"))
            else:
                changes.append((row, sc, old, status, f"状态 {old_s} → {status}"))
        else:
            note = f"状态 {old_s or '(空)'} → {status}"
            if old_s and old_s not in STATUS_RANK:
                note += f"（原值「{old_s}」是旧口径，本行按新口径改写，其余行不动）"
            changes.append((row, sc, old, status, note))

        if job.get("number") is not None:
            if cols["number"] is None:
                raise TableError("找不到「编号」列，拒绝写入")
            nc = cols["number"]
            old_n = ws.cell(row, nc).value
            if str(old_n) == str(job["number"]):
                changes.append((row, nc, old_n, old_n, "编号已一致，unchanged"))
            else:
                changes.append((row, nc, old_n, job["number"], "写编号"))

    # 2) 媒体列账号
    if job.get("account"):
        mc = media_col(cols, job["account"], job.get("account_column"))
        old_a = ws.cell(row, mc).value
        old_as = "" if old_a is None else str(old_a).strip()
        if old_as == job["account"]:
            changes.append((row, mc, old_a, old_as, "媒体列已一致，unchanged"))
        else:
            changes.append((row, mc, old_a, job["account"], f"媒体列写入账号 {job['account']}"))

    # 3) 自媒体发布链接（末尾「发布链接」列）
    if job.get("link"):
        write_link(ws, row, ensure_col(ws, cols, "link", LINK_HEADER, changes, cursor),
                   job["link"], "发布链接", changes)

    # 4) 网页链接（末尾「网页链接」列）
    if job.get("site_link"):
        write_link(ws, row, ensure_col(ws, cols, "site_link", SITE_LINK_HEADER, changes, cursor),
                   job["site_link"], "网页链接", changes)

    if not changes:
        raise TableError("这个任务没有任何可写项（至少要给 status / account / link / site_link 之一）")
    return changes


def apply(xlsx: Path, jobs: list, dry_run: bool) -> int:
    if not xlsx.exists():
        raise TableError(f"表格不存在：{xlsx}")

    # 先只读一遍，做全部校验（任一失败则整批不落盘）
    wb = openpyxl.load_workbook(xlsx)          # 不写 data_only：避免保存时丢公式
    ws = wb[wb.sheetnames[0]]
    cols = find_columns(ws)

    out(f"表格    : {xlsx}")
    out(f"工作表  : {wb.sheetnames[0]}  （数据区 2..{ws.max_row} 行，共 {ws.max_column} 列）")
    out(f"列定位  : 编号={cols['number']} 状态={cols['status']} 域名={cols['domain']} "
        f"发布链接={cols['link'] or '(将新增)'} 网页链接={cols['site_link'] or '(将新增)'}")
    out(f"表头实值: {[f'{openpyxl.utils.get_column_letter(c)}:{h}' for c, h in cols['headers'].items()]}")
    out()

    cursor = [ws.max_column + 1]          # 下一个可用的新列号（连加两列时自己往前推）
    all_changes = []
    for job in jobs:
        row = locate_row(ws, cols, domain=job.get("domain"), number=job.get("number"), row=job.get("row"))
        dom = ws.cell(row, cols["domain"]).value
        out(f"→ 命中的行：第 {row} 行（域名={dom!r}）")
        all_changes += plan_one(ws, cols, row, job, cursor)

    # 打印变更对照
    out()
    out(f"{'单元格':<8} {'原值':<28} {'新值':<28} 说明")
    out("-" * 100)
    changed = 0
    for r, c, old, new, why in all_changes:
        ref = f"{openpyxl.utils.get_column_letter(c)}{r}"
        o = "" if old is None else str(old)
        n = "" if new is None else str(new)
        if "unchanged" in why:
            out(f"{ref:<8} {o[:26]:<28} {n[:26]:<28} {why}")
        else:
            changed += 1
            out(f"{ref:<8} {o[:26]:<28} {n[:26]:<28} {why}  ← 改")
    out("-" * 100)
    out(f"实际需改动 {changed} 个单元格；{len(all_changes) - changed} 个已是目标值。")

    if changed == 0:
        out("\n无差异，不落盘、不备份。")
        return 0

    if dry_run:
        out("\n[dry-run] 未写盘、未备份。去掉 --dry-run 才真正执行。")
        return 0

    # 真写：先备份
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup = xlsx.with_name(f"{xlsx.stem}-备份-{stamp}{xlsx.suffix}")
    shutil.copy2(xlsx, backup)
    out(f"\n已备份 → {backup}")

    for r, c, _old, new, why in all_changes:
        if "unchanged" in why:
            continue
        ws.cell(row=r, column=c).value = new
    wb.save(xlsx)
    out(f"已写入 → {xlsx}")

    # 回读验证
    wb2 = openpyxl.load_workbook(xlsx)
    ws2 = wb2[wb2.sheetnames[0]]
    out("\n回读验证：")
    for r, c, _old, new, why in all_changes:
        if "unchanged" in why:
            continue
        got = ws2.cell(r, c).value
        ref = f"{openpyxl.utils.get_column_letter(c)}{r}"
        flag = "ok" if str(got) == str(new) else "MISMATCH"
        out(f"  {ref} = {got!r}  [{flag}]")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description="来源表格回写（状态机 + 白名单，详见 references/table-writeback.md）")
    ap.add_argument("--xlsx", required=True, type=Path)
    ap.add_argument("--domain", help="按「域名」列定位目标行")
    ap.add_argument("--number", type=int, help="按「编号」列定位目标行")
    ap.add_argument("--row", type=int, help="直接指定行号（会回读域名校验）")
    ap.add_argument("--status", choices=VALID_STATUS, help="写入状态：已写作 / 已发布")
    ap.add_argument("--account", help="发布账号代号，写入表头等于该代号的媒体列")
    ap.add_argument("--account-column", help="媒体列表头（默认与 --account 相同）")
    ap.add_argument("--link", help="自媒体发布链接，写入末尾「发布链接」列（不存在则新增一列）")
    ap.add_argument("--site-link", dest="site_link",
                    help="上传到 FTP 后的网页链接，写入末尾「网页链接」列（不存在则新增一列）")
    ap.add_argument("--batch", type=Path, help="批量任务 JSON（形状见 references/table-writeback.md）")
    ap.add_argument("--dry-run", action="store_true", help="只打印将要改动的单元格，不落盘")
    a = ap.parse_args()

    try:
        if a.batch:
            jobs = json.loads(a.batch.read_text(encoding="utf-8"))
            if not isinstance(jobs, list) or not jobs:
                raise TableError("--batch 文件必须是含至少一个对象的 JSON 数组")
        else:
            jobs = [{"domain": a.domain, "number": a.number, "row": a.row,
                     "status": a.status, "account": a.account,
                     "account_column": a.account_column, "link": a.link,
                     "site_link": a.site_link}]
        return apply(a.xlsx, jobs, a.dry_run)
    except TableError as e:
        out(f"拒绝写入：{e}")
        return 2
    except Exception as e:  # noqa: BLE001
        out(f"执行失败：{type(e).__name__}: {e}")
        return 1


if __name__ == "__main__":
    sys.exit(main())
